from pathlib import Path

import pymupdf
from pptx import Presentation


MAX_EXTRACTION_UNITS = 500
MAX_EXTRACTED_CHARACTERS = 2_000_000


def extract_pdf(path: Path) -> list[dict[str, int | str | bool]]:
    pages: list[dict[str, int | str | bool]] = []
    character_count = 0

    with pymupdf.open(path) as document:
        if document.page_count > MAX_EXTRACTION_UNITS:
            raise ValueError(
                f"This PDF has too many pages. The maximum is {MAX_EXTRACTION_UNITS} pages."
            )

        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            character_count += len(text)
            if character_count > MAX_EXTRACTED_CHARACTERS:
                raise ValueError(
                    "This document contains too much text to extract in one request."
                )
            pages.append(
                {
                    "reference_type": "page",
                    "reference_number": page_number,
                    "text": text,
                    "has_text": bool(text),
                }
            )

    return pages


def extract_pptx(path: Path) -> list[dict[str, int | str | bool]]:
    slides: list[dict[str, int | str | bool]] = []
    character_count = 0
    presentation = Presentation(path)

    if len(presentation.slides) > MAX_EXTRACTION_UNITS:
        raise ValueError(
            f"This presentation has too many slides. The maximum is {MAX_EXTRACTION_UNITS} slides."
        )

    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_parts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        text = "\n".join(text_parts)
        character_count += len(text)
        if character_count > MAX_EXTRACTED_CHARACTERS:
            raise ValueError(
                "This presentation contains too much text to extract in one request."
            )
        slides.append(
            {
                "reference_type": "slide",
                "reference_number": slide_number,
                "text": text,
                "has_text": bool(text),
            }
        )

    return slides
